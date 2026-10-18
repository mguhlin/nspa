"""Shared layouts for the HTML deck and editable PowerPoint slide revisions."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
from html import escape
import base64,re
from lxml import html,etree as ET
from PIL import Image
from activities import PARAGRAPH,RUBRIC,presenter_notes,ACTIVITIES
from demo_one import SOURCE,DRAFT

ROOT=Path(__file__).resolve().parents[2]
INK='#2E4F66';TEAL='#007681';ORANGE='#FF7233'
def t(words,x,y,w,h,size=28,bold=False,color=INK,heading=False):
 return dict(kind='text',words=words,x=x,y=y,w=w,h=h,size=size,bold=bold,color=color,heading=heading)
def image(path,x,y,w,h,alt):return dict(kind='image',path=path,x=x,y=y,w=w,h=h,alt=alt)
def table(rows,x,y,widths,heights,size=23):return dict(kind='table',rows=rows,x=x,y=y,w=sum(widths),h=sum(heights),widths=widths,heights=heights,size=size)
def layout(title,subtitle,action,blocks):
 return [t(title,75,42,1130,64,40,True,ORANGE,True),t(subtitle,75,106,1130,54,24)]+blocks+[t(action,75,583,1130,58,23,True)]

def layouts():
 out={}
 out[3]=layout('Our two perspectives on the same review','I need a process my team and applicants can understand.',
 'I can explain where AI helps and who is responsible for decisions.',[
 image('revised/two-perspectives.webp',75,184,555,333,'A review coordinator and an applicant discuss the same application with equal dignity.'),
 t('My scholarship team needs',680,188,525,42,29,True,TEAL),t('A manageable workload.\nConsistent application of the rules.\nEvidence we can check and explain.',680,243,525,120,28),
 t('From an applicant’s perspective',680,378,525,42,29,True,TEAL),t('I understand the rules.\nMy story is represented accurately.\nI can ask for a correction.',680,433,525,112,28)])
 labels=['Role','Task','Context','Criteria','Format','Constraints']
 examples=['Assist our\ncoordinator','Check recorded\ndocuments','Portal and\napproved inbox','Use our\nrequirements','Evidence and\nstaff-check table','Staff verify\nand decide']
 out[6]=layout('My six-part review request','Each part gives the assistant information it would otherwise have to guess.',
 'My companion includes a worked example of every part and a complete request to copy.',
 [image('revised/six-prompt-parts.webp',75,178,1130,240,'Six symbols represent role, task, context, criteria, format and constraints.')]+[item for i,(label,example) in enumerate(zip(labels,examples)) for item in [t(label,75+i*190,426,180,38,27,True,TEAL),t(example,75+i*190,471,180,80,23)]])
 out[7]=layout('The score descriptions I am using','Community impact: I connect my rating to the application’s words.',
 'If another reviewer and I interpret a score description differently, we compare our evidence and clarify the wording.',[
 image('conversation/perspectives.png',75,197,435,326,'Two reviewers inspect the same source through different perspectives.'),
 table([['Level','Evidence described']]+[[level,words] for level,words in RUBRIC],550,177,[170,485],[46,78,78,78,78],22)])
 out[10]=layout('Demo 1 · Our document follow-up plan','Our team has 40 flags and one hour before reviewer assignments begin.',
 'Our plan names who verifies, when to contact an applicant, and how to resolve an unfinished check.',[
 table([['A-102 · P1, portal record','Prepared AI draft'],[SOURCE,DRAFT]],75,174,[565,565],[42,88],25),
 image('revised/document-follow-up.webp',75,311,1130,211,'Portal and inbox records flow to staff verification before an applicant is contacted.'),
 t('Portal + approved inbox',75,531,355,32,24,True,TEAL),t('Staff verification',470,531,355,32,24,True,TEAL),t('Applicant contact',875,531,330,32,24,True,TEAL)])
 out[11]=layout('Demo 2 · My rating and the words behind it','I compare one application paragraph with the score descriptions on slide 7.',
 'My explanation connects exact words to a score description and names what I still do not know.',[
 t('Application paragraph · E1',75,181,535,38,26,True,TEAL),t('“'+PARAGRAPH+'”',75,230,535,320,28),
 image('photorealistic/pair-review.png',665,178,540,270,'Two reviewers compare the source material and discuss their reasons.'),
 t('My rating: __ / 4',690,462,515,45,34,True,TEAL),t('My evidence: “__”',690,516,515,40,28)])
 blocks=[]
 records=[('A-101 · T1','“I care for my younger brother after school.”','Caregiving','An explicit care responsibility'),('A-102 · T1','“I plan to complete a nursing certificate.”','Healthcare goal','A stated study or work goal'),('A-103 · T1','“Neither parent attended college.”','Parent-education context','Eligibility definition still needed')]
 for i,(ident,words,tag,limit) in enumerate(records):
  y=183+i*124
  blocks += [t(ident,75,y,650,32,22,True,TEAL),t(words,75,y+37,650,70,28),t('→',765,y+33,100,65,50,True,TEAL),t(tag,875,y+13,330,68,27,True,TEAL),t(limit,875,y+84,330,45,21)]
 out[12]=layout('Demo 3 · The evidence behind my summary','A tag describes supplied words. It does not establish character or eligibility.',
 'My neutral summary preserves the applicant’s account and makes uncertainty visible.',blocks)
 blocks=[]
 for i,(large,label,words) in enumerate([('3 runs','Same record and prompt','I compare flags, quotations, omissions, and uncertainty.'),('1 change','Only a name changes','I investigate an unexplained difference in wording or judgment.'),('Staff comparison','Independent judgments','Our team checks drafts against staff review and the original records.')]):
  x=75+i*385
  blocks += [t(large,x,205,360,85,51 if i<2 else 39,True,TEAL),t(label,x,311,360,65,28,True),t(words,x,400,340,135,27)]
 out[13]=layout('Our first tests before a pilot','I can look for errors and inconsistent behavior before relying on a draft.',
 'A small practice set can reveal a problem. Our team still needs broader evaluation.',blocks)
 blocks=[image('revised/four-policy-choices.webp',75,172,1130,253,'Four equally weighted policy pathways illustrate prohibit, disclose, bounded use, and embrace.')]
 for i,(label,words) in enumerate([('Prohibit','No generative AI help.\nDefine exceptions.'),('Disclose','Permitted help is\nbriefly explained.'),('Bounded use','Allowed assistance\nhas clear limits.'),('Embrace','Tools are allowed.\nClaims stay truthful.')]):
  x=75+i*288
  blocks += [t(label,x,438,265,44,30,True,TEAL),t(words,x,491,265,66,24)]
 out[20]=layout('The applicant AI rule my organization chooses','Each approach needs a clear purpose and examples applicants can understand.',
 'My team publishes the rule before applications open and applies it consistently.',blocks)
 out[22]=layout('The three gaps I will work on first','I can use the full checklist in my companion to choose a manageable priority.',
 'My plan names an owner, evidence of progress, a review date, and a pause condition.',[
 image('infographics/readiness.png',75,177,1130,230,'Learning, recording a plan, and checking progress form a connected route.'),
 t('Tools and data',75,427,340,45,29,True,TEAL),t('Approved service.\nMinimum necessary input.',75,482,340,78,25),
 t('Evidence and review',470,427,350,45,29,True,TEAL),t('Clear criteria.\nTraceable, checked drafts.',470,482,350,78,25),
 t('Rules and transparency',865,427,340,45,28,True,TEAL),t('Applicant guidance.\nOur use and correction route.',865,482,340,78,25)])
 return out

def updated_notes(number,old):
 # A single introductory disclosure replaces repeated warnings in every slide.
 paras=[p.text_content().strip() for p in old.xpath('./p') if not p.text_content().startswith(('All practice data and displayed demo responses','Source artwork is AI-generated; people are fictional.'))]
 if number==1:
  paras=['Our session runs from 3:30 to 5:00 PM. The route integrates the slides and eight activities. All cases, records, operational constraints, and prepared drafts are fictional teaching examples. No participant needs an AI account. Generated illustrations depict fictional situations.']+paras
  paras=[re.sub(r'Our 90-minute route integrates explanations and eight complete activities\. The earlier section cues were:[\s\S]*?The four backup slides sit outside that route\.','Our timed route includes explanations, eight activities and transitions. The four backup slides are optional.',x) for x in paras]
 extra=presenter_notes(number)
 if extra:paras=extra
 elif number in layouts():
  paras= ['This visual supports the matching activity in the companion. Read the slide’s first-person statement as the participant’s learning target. The companion contains the rules, source material, examples and response space.']
 if number==7:paras=['The term reviewers means the scholarship team applying the rubric, including participants and their partners during the activity. Read the four descriptions before assigning a rating. If two reviewers interpret a score description differently, ask each to explain the exact words behind the judgment, then clarify the adopted rubric before real review. The repair-table paragraph on slide 11 is the material used in this activity; there is no missing longer essay.']
 if number==3:paras=['Use this visual to compare needs, not to ask an unexplained question. A scholarship team needs a manageable workload, consistent rules and evidence it can explain. From an applicant’s perspective, the needs include knowing the rules, accurate representation and a correction route. Ask participants to use the Trust activity in the companion to propose one improvement and explain whose need it addresses.']
 if number==20:paras=['The four illustrated approaches are policy choices, not a ranking of ethical quality. A prohibition needs scope and accessibility exceptions. Disclosure needs a short, usable statement and an explanation of how it affects review. Bounded use needs permitted and prohibited examples. Embrace still requires truthful claims and examination of substance. The supplied practice rule in the companion combines bounded use with disclosure.']
 if number==22:paras=['The complete twelve-item checklist is in the companion. The slide groups it into tools and data, evidence and review, and rules and transparency. Ask participants to choose three gaps without promising to fix everything at once. The final activity turns a priority into a 30 / 60 / 90-day plan.']
 if number==24:paras=[re.sub(r'The session is 90 minutes,[\s\S]*?match this presentation\.','The complete activities and participant downloads are on the website.',x) for x in paras]
 return [re.sub(r'^\d+ minutes?(?:, [^.]+)?\.\s*','',x).replace('timed sixty-minute route','timed ninety-minute route') for x in paras]

def patch_html_deck(text):
 specs=layouts();index=0
 def replace(match):
  nonlocal index
  index+=1;slide=html.fromstring(match[0]);old=slide.xpath('./div[@class="notes"]')[0]
  originals={p.text_content().strip():html.tostring(p,encoding='unicode') for p in old.xpath('./p')}
  words=updated_notes(index,old)
  for child in list(old):old.remove(child)
  for paragraph in words:
   if paragraph in originals:old.append(html.fromstring(originals[paragraph]))
   else:ET.SubElement(old,'p').text=paragraph
  if index in specs:
   body=slide.xpath('./div[@class="slide-body"]')[0]
   for child in list(body):body.remove(child)
   body.text=None
   for block in specs[index]:
    if block['kind']=='text':
     node=ET.SubElement(body,'h2' if block['heading'] else 'p',{'class':'txt'+(' slide-title' if block['heading'] else ''),'style':f'left:{block["x"]}px;top:{block["y"]}px;width:{block["w"]}px;height:{block["h"]}px;font-size:{block["size"]}px;font-weight:{700 if block["bold"] else 400};color:{block["color"]};'})
     for i,line in enumerate(block['words'].split('\n')):
      if i:br=ET.SubElement(node,'br');br.tail=line
      else:node.text=line
    elif block['kind']=='image':
     data=base64.b64encode((ROOT/'2026/images'/block['path']).read_bytes()).decode();kind='webp' if block['path'].endswith('.webp') else 'png'
     ET.SubElement(body,'img',{'class':'art','src':f'data:image/{kind};base64,{data}','alt':block['alt'],'style':f'left:{block["x"]}px;top:{block["y"]}px;width:{block["w"]}px;height:{block["h"]}px;object-fit:contain;'})
    else:
     tab=ET.SubElement(body,'table',{'class':'native-table','style':f'left:{block["x"]}px;top:{block["y"]}px;width:{block["w"]}px;'})
     grid=ET.SubElement(tab,'colgroup')
     for width in block['widths']:ET.SubElement(grid,'col',{'style':f'width:{width}px'})
     for i,row in enumerate(block['rows']):
      tr=ET.SubElement(tab,'tr')
      for j,value in enumerate(row):ET.SubElement(tr,'th' if i==0 else 'td',{'data-label':block['rows'][0][j],'style':f'font-size:{block["size"]}px;height:{block["heights"][i]}px;background:{TEAL if i==0 else "#FFFFFF"};color:{"#FFFFFF" if i==0 else INK};'}).text=value
   ET.SubElement(body,'div',{'class':'rule','style':'left:75px;top:568px;width:1130px;border-color:#FFCD34;','aria-hidden':'true'})
   slide.set('aria-label',f'Slide {index}: '+specs[index][0]['words'])
  else:
   for node in slide.xpath('./div[@class="slide-body"]//*[text()]'):
    if node.text:node.text=node.text.replace('Read the full essay and rubric.','Read the paragraph and score descriptions.')
  return html.tostring(slide,encoding='unicode')
 return re.sub(r'<section\b[\s\S]*?</section>',replace,text)

def patch_powerpoint(path,html_slides):
 from pptx import Presentation
 from pptx.util import Pt
 from pptx.dml.color import RGBColor
 specs=layouts();raw=Path(path).read_bytes();prs=Presentation(BytesIO(raw))
 for number,blocks in specs.items():
  slide=prs.slides[number-1]
  for shape in list(slide.shapes):
   full_background=shape.shape_type==13 and shape.width>=12100000 and shape.height>=6800000
   footer=shape.top>=648*9525
   line=shape.top==568*9525 and shape.height==0
   if not (full_background or footer or line):slide.shapes._spTree.remove(shape._element)
  for block in blocks:
   box=[round(block[k]*9525) for k in ['x','y','w','h']]
   if block['kind']=='text':
    shape=slide.shapes.add_textbox(*box);frame=shape.text_frame;frame.clear();frame.word_wrap=True
    frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
    for i,line in enumerate(block['words'].split('\n')):
     para=frame.paragraphs[0] if i==0 else frame.add_paragraph();para.text=line;para.space_before=para.space_after=Pt(0);para.line_spacing=1.2
     for run in para.runs:run.font.name='Liberation Sans';run.font.size=Pt(block['size']*.75);run.font.bold=block['bold'];run.font.color.rgb=RGBColor.from_string(block['color'][1:])
   elif block['kind']=='image':
    source=ROOT/'2026/images'/block['path'];im=Image.open(source);iw,ih=im.size
    # Use PNG for portable editable PowerPoint picture parts.
    buf=BytesIO();im.convert('RGB').save(buf,format='PNG');buf.seek(0)
    scale=min(box[2]/iw,box[3]/ih);w=round(iw*scale);h=round(ih*scale)
    slide.shapes.add_picture(buf,box[0]+(box[2]-w)//2,box[1]+(box[3]-h)//2,width=w,height=h)
   else:
    shape=slide.shapes.add_table(len(block['rows']),len(block['rows'][0]),*box);tab=shape.table
    for j,width in enumerate(block['widths']):tab.columns[j].width=round(width*9525)
    for i,row in enumerate(block['rows']):
     tab.rows[i].height=round(block['heights'][i]*9525)
     for j,words in enumerate(row):
      cell=tab.cell(i,j);cell.text=words;cell.margin_left=cell.margin_right=round(10*9525);cell.margin_top=cell.margin_bottom=round(6*9525)
      cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('007681' if i==0 else 'FFFFFF')
      for para in cell.text_frame.paragraphs:
       para.space_before=para.space_after=Pt(0);para.line_spacing=1.15
       for run in para.runs:run.font.name='Liberation Sans';run.font.size=Pt(block['size']*.75);run.font.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '2E4F66')
 out=BytesIO();prs.save(out)
 with ZipFile(BytesIO(raw)) as old,ZipFile(BytesIO(out.getvalue())) as new:
  parts={i.filename:old.read(i.filename) for i in old.infolist()}
  for name in new.namelist():
   match=re.fullmatch(r'ppt/slides/(?:_rels/)?slide(\d+)\.xml(?:\.rels)?',name)
   if (match and int(match[1]) in specs) or name not in parts:parts[name]=new.read(name)
  # Notes are taken from the same final HTML source as the script and guide.
  ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
  for i,slide in enumerate(html_slides,1):
   name=f'ppt/notesSlides/notesSlide{i}.xml';root=ET.fromstring(parts[name])
   body=root.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]/p:txBody',namespaces=ns)
   if not body:continue
   body=body[0]
   for para in body.findall('a:p',ns):body.remove(para)
   for para in slide.xpath('./div[@class="notes"]/p'):
    p=ET.SubElement(body,'{'+ns['a']+'}p');r=ET.SubElement(p,'{'+ns['a']+'}r');ET.SubElement(r,'{'+ns['a']+'}t').text=para.text_content()
   parts[name]=ET.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
 with ZipFile(path,'w',ZIP_DEFLATED) as final:
  for name,data in parts.items():final.writestr(name,data)
